import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation('TemporalAML_review_1(Anshu).pptx')
s4 = prs.slides[3]

DARK_NAVY = RGBColor(0, 51, 102)     # Header bold
DARK_CHARCOAL = RGBColor(35, 35, 35) # Body text
MUTED_GRAY = RGBColor(70, 70, 70)    # Sub-bullets
BLUE_TITLE = RGBColor(0, 51, 153)    # Slide Title

# Find title and body shapes
t4 = None
b4 = None
for s in s4.shapes:
    if s.name == 'TextBox 11':
        t4 = s
    elif s.name == 'TextBox 12':
        b4 = s

if t4:
    t4.text_frame.clear()
    p = t4.text_frame.paragraphs[0]
    p.text = "Introduction: Cryptocurrency, Bitcoin & AML Foundations"
    p.runs[0].font.size = Pt(22)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = BLUE_TITLE

if b4:
    b4.top = Inches(1.6)
    b4.height = Inches(5.4)
    tf = b4.text_frame
    tf.word_wrap = True
    tf.clear()

    items = [
        ("1. What is Cryptocurrency?", "A decentralized, peer-to-peer digital monetary medium operating on immutable distributed ledgers (Blockchains) without central bank intermediaries, solving the double-spending problem via cryptographic consensus (PoW/PoS)."),
        ("2. What is Bitcoin & The UTXO Model?", "The first decentralized cryptocurrency (Satoshi Nakamoto, 2008) with a hard cap of 21 Million BTC; utilizes the Unspent Transaction Output (UTXO) model where transactions consume and produce outputs, naturally forming a Directed Acyclic Graph (DAG)."),
        ("3. What is a Digital Signature?", "An asymmetric cryptographic mechanism (ECDSA over secp256k1 curve) providing transaction Authentication, Integrity, and Non-Repudiation; transactions are signed with a private key and verified globally via the sender's public key (wallet address)."),
        ("4. Money Laundering & Pseudonymity Paradox:", "Converting illicit proceeds into legitimate funds ($800B to $2 Trillion annually, 2–5% of global GDP); while blockchains are publicly auditable, wallet addresses are pseudonymous hashes, enabling illicit syndicates to execute borderless transfers without verified identities."),
        ("5. Laundering Typologies in Crypto Graphs:", "Illicit entities obscure fund origins using 3 specific structural graph patterns:"),
        [
            "• Circular Transfers: Closed loops (A→B→C→A) designed for wash-trading, spoofing volumes, and confusing audit algorithms.",
            "• Layering Chains: Rapid sequential transfers across consecutive timestamps to distance dirty funds from origin crimes.",
            "• Smurfing (Structuring): Splitting large sums into micro-transactions below regulatory reporting thresholds ($10,000 limit)."
        ],
        ("6. Proposed Solution (TemporalAML):", "A continuous-time Temporal Graph Attention Network (TGAT) with Learnable Fourier Time Encoding that detects all 3 typologies simultaneously and generates auditable causal subgraph evidence.")
    ]

    for i, item in enumerate(items):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.space_after = Pt(3)
        p.space_before = Pt(1)

        if isinstance(item, tuple):
            head, body = item
            r1 = p.add_run()
            r1.text = head + " " if head else ""
            r1.font.bold = True
            r1.font.size = Pt(10.5)
            r1.font.color.rgb = DARK_NAVY

            if body:
                r2 = p.add_run()
                r2.text = body
                r2.font.bold = False
                r2.font.size = Pt(10.5)
                r2.font.color.rgb = DARK_CHARCOAL
        elif isinstance(item, list):
            for sub in item:
                sub_p = tf.add_paragraph()
                sub_p.level = 1
                sub_p.space_after = Pt(1)
                sub_p.space_before = Pt(1)
                r = sub_p.add_run()
                r.text = sub
                r.font.size = Pt(9.5)
                r.font.color.rgb = MUTED_GRAY

prs.save('TemporalAML_review_1(Anshu).pptx')
print("Successfully updated Slide 4 of TemporalAML_review_1(Anshu).pptx!")
